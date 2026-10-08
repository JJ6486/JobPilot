"""The tracker API end to end: analyze -> review -> edit -> status -> analytics, with a fake LLM."""

import shutil
from datetime import UTC, datetime, timedelta
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from app.api.deps import get_provider
from app.config import Settings, get_settings
from app.db import session as db_session
from app.llm.schemas import AnalysisResult, TailorResult
from app.main import app
from app.tasks import TaskRegistry, get_registry
from tests.conftest import FIXTURES, JDS, FakeProvider, make_analysis, make_tailor


@pytest.fixture
def api(settings: Settings, tmp_path: Path):
    shutil.copy(FIXTURES / "truth_bank.yaml", settings.truth_bank_path)
    provider = FakeProvider({AnalysisResult: make_analysis(), TailorResult: make_tailor()})
    app.dependency_overrides[get_settings] = lambda: settings
    app.dependency_overrides[get_provider] = lambda: provider
    registry = TaskRegistry(sync=True)
    app.dependency_overrides[get_registry] = lambda: registry
    with TestClient(app) as client:  # runs the lifespan: migrations on the temp database
        client.provider = provider  # type: ignore[attr-defined]
        yield client
    app.dependency_overrides.clear()


def _analyze(api: TestClient, **extra) -> int:
    body = {
        "jd_text": (JDS / "software_ict.txt").read_text("utf-8"),
        "source_portal": "Indeed",
        **extra,
    }
    res = api.post("/api/analyze", json=body)
    assert res.status_code == 202
    task = api.get(f"/api/analyze/{res.json()['task_id']}").json()
    assert task["state"] == "done", task
    assert [s["step"] for s in task["steps"]] == [
        "parse",
        "dedupe",
        "score",
        "tailor",
        "letter",
        "save",
    ]
    return task["application_id"]


def test_analyze_then_review(api: TestClient) -> None:
    app_id = _analyze(api)
    detail = api.get(f"/api/applications/{app_id}").json()
    assert detail["status"] == "tailored"
    assert detail["fit_score"] == 85 and detail["recommendation"] == "apply"
    assert detail["job"]["company"] == "Example Corp"
    assert "Himmah" in detail["job"]["jd_text"]  # full JD snapshot kept
    assert detail["resume"]["experience"][0]["source_id"] == "exp_1"
    assert detail["cover_letter"].endswith("Test Candidate")
    assert detail["has_docx"] and not detail["download_blocked"]
    assert [e["type"] for e in detail["events"]] == ["created", "status_change"]
    res = api.get(f"/api/applications/{app_id}/resume.docx")
    assert res.status_code == 200 and res.content[:2] == b"PK"


def test_applied_sets_follow_up_and_logs_events(api: TestClient) -> None:
    app_id = _analyze(api)
    detail = api.post(
        f"/api/applications/{app_id}/status", json={"status": "applied", "note": "via Indeed"}
    ).json()
    applied = datetime.fromisoformat(detail["applied_at"])
    assert datetime.fromisoformat(detail["follow_up_at"]) - applied == timedelta(days=7)
    assert detail["events"][-1] | {"id": 0, "created_at": ""} == {
        "id": 0,
        "type": "status_change",
        "from_status": "tailored",
        "to_status": "applied",
        "note": "via Indeed",
        "created_at": "",
    }
    assert (
        api.post(f"/api/applications/{app_id}/status", json={"status": "bogus"}).status_code == 422
    )


def test_follow_up_due_and_ghost_suggestion(api: TestClient) -> None:
    app_id = _analyze(api)
    api.post(f"/api/applications/{app_id}/status", json={"status": "applied"})
    past = (datetime.now(UTC) - timedelta(days=30)).isoformat()
    api.patch(f"/api/applications/{app_id}", json={"applied_at": past, "follow_up_at": past})
    listing = api.get("/api/applications").json()
    row = listing["items"][0]
    assert row["follow_up_due"] and row["suggest_ghosted"]
    assert listing["follow_ups_due"] == 1 and listing["ghost_suggestions"] == 1
    api.patch(f"/api/applications/{app_id}", json={"ghost_dismissed": True})  # user dismisses
    assert api.get("/api/applications").json()["ghost_suggestions"] == 0
    assert (
        api.get(f"/api/applications/{app_id}").json()["status"] == "applied"
    )  # never auto-applied


def test_edit_with_unsupported_claim_blocks_download_until_confirmed(api: TestClient) -> None:
    app_id = _analyze(api)
    resume = api.get(f"/api/applications/{app_id}").json()["resume"]
    resume["summary"] = "Engineer experienced with React and Kubernetes."
    detail = api.patch(f"/api/applications/{app_id}", json={"resume": resume}).json()
    assert detail["high_warnings"] >= 1 and detail["download_blocked"]
    api.post(f"/api/applications/{app_id}/rerender")
    assert api.get(f"/api/applications/{app_id}/resume.docx").status_code == 409
    detail = api.patch(f"/api/applications/{app_id}", json={"warnings_confirmed": True}).json()
    assert not detail["download_blocked"]
    assert api.get(f"/api/applications/{app_id}/resume.docx").status_code == 200
    assert "edited" in [e["type"] for e in detail["events"]]


def test_rerender_uses_edited_content_without_llm(api: TestClient) -> None:
    from docx import Document

    app_id = _analyze(api)
    calls_before = len(api.provider.calls)  # type: ignore[attr-defined]
    resume = api.get(f"/api/applications/{app_id}").json()["resume"]
    resume["summary"] = "Software Engineer who builds React apps."
    api.patch(f"/api/applications/{app_id}", json={"resume": resume})
    api.post(f"/api/applications/{app_id}/rerender")
    path = Path(
        api.get(f"/api/applications/{app_id}/resume.docx")
        .headers["content-disposition"]
        .split('"')[1]
    )
    stored = db_docx_path(api, app_id)
    assert path.name == Path(stored).name
    assert "Software Engineer who builds React apps." in "\n".join(
        p.text for p in Document(stored).paragraphs
    )
    assert len(api.provider.calls) == calls_before  # type: ignore[attr-defined]


def db_docx_path(api: TestClient, app_id: int) -> str:
    from sqlmodel import Session

    from app.db.models import Application

    settings = app.dependency_overrides[get_settings]()
    with Session(db_session.get_engine(settings)) as s:
        return s.get(Application, app_id).resume_docx_path  # type: ignore[union-attr,return-value]


def test_regenerate_cover_letter(api: TestClient) -> None:
    app_id = _analyze(api)
    api.provider.responses[TailorResult] = make_tailor(  # type: ignore[attr-defined]
        whatsapp_message="Hello again, I applied for the Software Engineer role."
    )
    res = api.post(f"/api/applications/{app_id}/regenerate", json={"target": "cover_letter"})
    assert res.status_code == 202
    detail = api.get(f"/api/applications/{app_id}").json()
    assert detail["whatsapp_message"].startswith("Hello again")
    assert detail["events"][-1]["type"] == "regenerated"


def test_data_persists_across_restarts(api: TestClient, settings: Settings) -> None:
    app_id = _analyze(api)
    api.post(f"/api/applications/{app_id}/status", json={"status": "applied"})
    db_session._engine_for.cache_clear()  # drop the connection pool, like a process restart
    with TestClient(app) as fresh:
        detail = fresh.get(f"/api/applications/{app_id}").json()
    assert detail["status"] == "applied" and detail["job"]["company"] == "Example Corp"


def test_filters_csv_and_duplicates(api: TestClient) -> None:
    first = _analyze(api)
    second = _analyze(api)
    dupes = api.get(f"/api/applications/{second}").json()["duplicates"]
    assert dupes and dupes[0]["job_id"] == str(first)
    api.post(f"/api/applications/{first}/status", json={"status": "applied"})
    assert [r["id"] for r in api.get("/api/applications?status=applied").json()["items"]] == [first]
    assert api.get("/api/applications?score_min=90").json()["total"] == 0
    assert api.get("/api/applications?q=example").json()["total"] == 2
    csv = api.get("/api/applications/export.csv").text
    assert csv.lstrip("﻿").startswith("id,created_at,status") and csv.count("\n") == 3


def test_analytics(api: TestClient) -> None:
    a, b = _analyze(api), _analyze(api)
    for app_id, status in ((a, "applied"), (b, "applied"), (b, "interview")):
        api.post(f"/api/applications/{app_id}/status", json={"status": status})
    summary = api.get("/api/analytics/summary").json()
    assert summary["funnel"] == {"applied": 2, "replied": 1, "interview": 1, "offer": 0}
    assert summary["response_by_portal"][0] == {
        "key": "indeed",
        "applied": 2,
        "replied": 1,
        "rate": 0.5,
    }
    assert summary["response_by_score_band"][0]["key"] == "70-100"
    assert summary["median_days_to_first_reply"] is not None
    gaps = api.get("/api/analytics/skill-gaps").json()
    assert gaps["missing"][0] == {"skill": "Kubernetes", "count": 2}
    assert gaps["adjacent"][0]["skill"] == "Vue"


def test_truth_bank_read_and_validated_write(api: TestClient, settings: Settings) -> None:
    doc = api.get("/api/truth-bank").json()
    assert doc["valid"] and "Test Candidate" in doc["yaml"]
    bad = doc["yaml"].replace("level: core", "level: expert", 1)
    res = api.put("/api/truth-bank", json={"yaml": bad})
    assert res.status_code == 422 and res.json()["detail"]["problems"]
    assert "level: expert" not in settings.truth_bank_path.read_text("utf-8")  # nothing written
    good = doc["yaml"].replace('tone: "professional, warm, concise"', 'tone: "friendly"')
    assert api.put("/api/truth-bank", json={"yaml": good}).json()["valid"]
    assert 'tone: "friendly"' in settings.truth_bank_path.read_text("utf-8")
    assert len(list((settings.data_dir / "backups").glob("truth_bank-*.yaml"))) == 1


def test_base_resumes_and_delete(api: TestClient, settings: Settings) -> None:
    families = {r["role_family"]: r for r in api.get("/api/base-resumes").json()}
    assert families["software_ict"]["template_found"]
    app_id = _analyze(api)
    assert (settings.outputs_dir / str(app_id)).is_dir()
    assert api.delete(f"/api/applications/{app_id}").status_code == 204
    assert api.get(f"/api/applications/{app_id}").status_code == 404
    assert not (settings.outputs_dir / str(app_id)).exists()


def test_truth_bank_validate_does_not_write(api: TestClient, settings: Settings) -> None:
    before = settings.truth_bank_path.read_text("utf-8")
    res = api.post("/api/truth-bank/validate", json={"yaml": "profile: [broken"}).json()
    assert not res["valid"] and res["problems"]
    assert settings.truth_bank_path.read_text("utf-8") == before


def test_base_resume_text(api: TestClient, settings: Settings) -> None:
    from docx import Document

    settings.base_resumes_dir.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_paragraph("TEST CANDIDATE")
    doc.add_paragraph("")
    doc.add_paragraph("Built things.")
    doc.save(str(settings.base_resumes_dir / "sw.docx"))
    assert api.get("/api/base-resumes/software_ict/text").json()["lines"] == [
        "TEST CANDIDATE",
        "Built things.",
    ]
    assert api.get("/api/base-resumes/nope/text").status_code == 404
